"""Generate tests/golden_files/merged_headers.xlsx for the ingestion tests.

The file has a merged title row above the real header row (a "merged-header"
layout). It is committed to the repo; this script only regenerates it.
Run: python scripts/gen_golden_xlsx.py
"""
from openpyxl import Workbook

HEADERS = ["product_name", "sale_timestamp", "quantity", "unit_price",
           "total_amount", "payment_method"]
ROWS = [
    ["Paracetamol", "2026-09-01", 2, 1.50, 3.00, "card"],
    ["Ibuprofen", "2026-09-01", 1, 2.00, 2.00, "cash"],
]

OUT = "tests/golden_files/merged_headers.xlsx"


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.merge_cells("A1:F1")
    ws["A1"] = "Pharmacy Sales Export"
    ws.append(HEADERS)
    for row in ROWS:
        ws.append(row)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()